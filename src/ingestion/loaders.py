"""Format-specific extraction. Every loader returns the same shape:

    {"layout": "document" | "slides",
     "units": [{"kind": "page" | "slide" | "section", "number": int | None, "title": str | None,
                "images": int, "lines": [line, ...]}],
     "page_count": int | None, "slide_count": int | None,
     "properties": {...embedded file properties, unverified...},
     "master_text": [...]}   # slide-master text, used only as metadata evidence

A line is {"text", "kind": heading|text|code|bullet|table|notes, "level", "para_break"} plus,
for PDFs, geometry ("top"/"bottom" as page fractions, "x0", "cw") used by the cleaner.
Structure is kept at line level; the cleaner assembles lines into blocks.
"""
import re
import struct
import zipfile
from collections import Counter
from pathlib import Path
from xml.etree import ElementTree as ET

MONO_FONT = re.compile(r"mono|monl|courier|consol|menlo|cmtt|typewriter|lucida console|source code", re.I)
BULLET_START = re.compile("^\\s*[•●○◦▪■□➢❖✓→➔\\-–—*·\uf0a7\uf0b7\uf0d8\uf076\uf0fc]\\s+")

# ponytail: regex code heuristic for documents without a monospace font (Word/PowerPoint exports).
# Misses code with no C punctuation; upgrade to a small classifier if it matters.
CODE_HINT = re.compile(
    r"(;\s*(//.*|/\*.*)?$)|(^[{}]\s*;?$)|(\{\s*(//.*)?$)|(^\}\s*(while|else)\b)|"
    r"(^(#\s*(define|include)\b|while\s*\(|if\s*\(|for\s*\(|do\s*\{|else\b|return\b|void\s+\w|int\s+\w|"
    r"struct\s+\w|typedef\b|semaphore\s+\w|pthread_(t|mutex_t|cond_t)\s+\w))|"
    r"((?<!:)//\s)|(^(//|/\*|\*/))|(^\w+(\[\w*\])?\s*(\+\+|--|[+\-*/%]?=(?!=))\s*[\w(\-][^=]*$)",
    re.I,
)


def looks_like_code(text):
    t = text.strip()
    if not t or BULLET_START.match(t) or re.match(r"^\(?[a-z0-9]{1,3}[.)]\s", t, re.I):
        return False
    words = t.split()
    if len(words) > 14 and ";" not in t and "{" not in t:
        return False
    if t.endswith(".") and len(words) >= 6:
        return False
    return bool(CODE_HINT.search(t))


def line(text, kind="text", level=0, para_break=False, **extra):
    return {"text": text, "kind": kind, "level": level, "para_break": para_break, **extra}


# ---------------------------------------------------------------- PDF
def load_pdf(path):
    import fitz

    doc = fitz.open(path)
    pages, size_chars, mono_chars, total_chars, landscape = [], Counter(), 0, 0, 0
    for page in doc:
        W, H = page.rect.width, page.rect.height
        landscape += W > H
        vlines = []  # visual lines: {"spans": [...], "y0", "y1"}
        seen = set()
        # TeX renders "_" as a short stroke inside a space glyph; recover it so pthread_cond_t stays intact
        strokes = [d["rect"] for d in page.get_drawings() if d["rect"].height < 1 and 0 < d["rect"].width < 12]
        for b in page.get_text("rawdict")["blocks"]:
            for l in b.get("lines", []):
                if abs(l["dir"][1]) > 0.1:  # rotated text (watermarks, margin stamps)
                    continue
                spans = []
                for s in l["spans"]:
                    t = _span_text(s, strokes)
                    if not t.strip():
                        continue
                    x0, y0, x1, y1 = s["bbox"]
                    key = (t.strip(), round(x0), round(y0))
                    if key in seen:  # identical text drawn twice (shadow / overprint)
                        continue
                    seen.add(key)
                    spans.append({"text": t, "x0": x0, "x1": x1, "y0": y0, "y1": y1, "size": s["size"],
                                  "bold": bool(s["flags"] & 16) or "bold" in s["font"].lower(),
                                  "mono": bool(s["flags"] & 8) or bool(MONO_FONT.search(s["font"]))})
                if spans:
                    _attach(vlines, spans)
        page_lines = []
        for v in vlines:
            v["spans"].sort(key=lambda s: s["x0"])
            big = max(s["size"] for s in v["spans"])
            while len(v["spans"]) > 1 and v["spans"][0]["text"].strip().isdigit() and v["spans"][0]["size"] < 0.75 * big:
                v["spans"].pop(0)  # code-listing gutter number glued to its code line
            text = _join_spans(v["spans"])
            if not text.strip():
                continue
            n = len(text.strip())
            sizes = Counter()
            for s in v["spans"]:
                sizes[round(s["size"], 1)] += len(s["text"].strip())
            size = sizes.most_common(1)[0][0]
            mono = sum(len(s["text"].strip()) for s in v["spans"] if s["mono"])
            mono_chars += mono
            total_chars += n
            size_chars[size] += n
            mono_spans = [s for s in v["spans"] if s["mono"] and s["text"].strip()]
            first = mono_spans[0] if mono_spans else v["spans"][0]
            page_lines.append({
                "text": text, "size": size, "mono_frac": mono / n,
                "bold": all(s["bold"] for s in v["spans"] if s["text"].strip()),
                "x0": round(first["x0"], 1), "x1": max(s["x1"] for s in v["spans"]),
                "y0": v["y0"], "y1": v["y1"], "top": round(v["y0"] / H, 3), "bottom": round(v["y1"] / H, 3),
                "cw": (first["x1"] - first["x0"]) / max(len(first["text"]), 1),
                "digits_only": text.strip().isdigit(),
            })
        pages.append({"lines": page_lines, "images": len(page.get_images()), "W": W})

    slides = landscape > len(doc) / 2
    body = max(size_chars.items(), key=lambda kv: kv[1])[0] if size_chars else 0
    font_code = total_chars and mono_chars / total_chars >= 0.02
    heading_sizes = sorted({s for s in size_chars if s >= body * 1.08}, reverse=True)
    units = []
    for no, p in enumerate(pages, 1):
        lines, title = [], None
        # slide decks printed to PDF: title = largest line(s) in the top quarter (anywhere on the title slide)
        title_size = max((l["size"] for l in p["lines"] if l["top"] < 0.25 or no == 1), default=0)
        right = max((l["x1"] for l in p["lines"]), default=0)
        prev = None
        for l in p["lines"]:
            if l["digits_only"] and l["size"] <= body * 0.7:
                continue  # code-listing gutter numbers and tiny slide numbers
            kind, level = "text", 0
            if slides:
                if l["size"] == title_size and title_size >= body * 1.15 and (l["top"] < 0.25 or no == 1):
                    kind, level = "heading", 1
            elif l["size"] >= body * 1.08 and l["mono_frac"] < 0.5:
                kind, level = "heading", min(heading_sizes.index(l["size"]) + 1, 3)
            elif l["bold"] and len(l["text"]) <= 80 and not re.match(r"^\s*(ans(wer)?|solution)\b", l["text"], re.I) \
                    and not l["text"].rstrip().endswith("."):
                kind, level = "heading", min(len(heading_sizes) + 1, 4)
            if kind == "text" and (l["mono_frac"] >= 0.7 if font_code else looks_like_code(l["text"])):
                kind = "code"
            brk = prev is None or _pdf_break(prev, l, right, slides)
            lines.append(line(l["text"], kind, level, brk, top=l["top"], bottom=l["bottom"], x0=l["x0"], cw=round(l["cw"], 2)))
            if kind == "heading" and title is None and slides:
                title = l["text"].strip()
            prev = l
        units.append({"kind": "page", "number": no, "title": title, "images": p["images"], "lines": lines})
    props = {k: v for k, v in (doc.metadata or {}).items() if v and k in ("title", "author", "creator", "producer", "creationDate")}
    return {"layout": "slides" if slides else "document", "units": units, "page_count": len(doc),
            "slide_count": None, "properties": props, "master_text": []}


def _span_text(span, strokes):
    out = []
    for ch in span["chars"]:
        c = ch["c"]
        if c == " " and strokes:
            x0, _, x1, y1 = ch["bbox"]
            if any(abs(r.y1 - y1) < 0.35 * span["size"] and min(r.x1, x1) - max(r.x0, x0) > 0.6 * (x1 - x0)
                   and r.width < 1.5 * (x1 - x0) for r in strokes):
                c = "_"
        out.append(c)
    return "".join(out)


def _attach(vlines, spans):
    """Merge a PyMuPDF line into an existing visual line when it sits on the same baseline right next to it
    (split headings, small caps). Side-by-side text boxes stay separate so reading order follows the boxes.
    Fragments overprinted on the same baseline are junk: the longest text wins."""
    y0, y1 = min(s["y0"] for s in spans), max(s["y1"] for s in spans)
    x0, x1 = min(s["x0"] for s in spans), max(s["x1"] for s in spans)
    size = max(s["size"] for s in spans)
    length = lambda sp: sum(len(s["text"].strip()) for s in sp)
    near, overlapped = [], []
    for v in vlines[-12:]:
        h = min(y1 - y0, v["y1"] - v["y0"])
        if abs((y0 + y1) / 2 - (v["y0"] + v["y1"]) / 2) > 0.5 * h:
            continue
        gap = max(x0 - max(s["x1"] for s in v["spans"]), min(s["x0"] for s in v["spans"]) - x1)
        (overlapped if gap < -2 else near).append((gap, v))
    if overlapped:
        if all(length(spans) > length(v["spans"]) for _, v in overlapped):
            vlines[vlines.index(overlapped[0][1])] = {"spans": spans, "y0": y0, "y1": y1}
            for _, v in overlapped[1:]:
                vlines.remove(v)
        return
    for gap, v in reversed(near):
        if gap < 1.5 * size:
            v["spans"] += spans
            v["y0"], v["y1"] = min(v["y0"], y0), max(v["y1"], y1)
            return
    vlines.append({"spans": spans, "y0": y0, "y1": y1})


def _join_spans(spans):
    out = ""
    for i, s in enumerate(spans):
        t = s["text"]
        if out and not out[-1].isspace() and not t[0].isspace() and s["x0"] - spans[i - 1]["x1"] > 0.15 * s["size"]:
            out += " "
        out += t
    return re.sub(r"\s+$", "", out)


def _pdf_break(prev, cur, right, slides):
    """True = new paragraph, "line" = hard line break inside a block, False = wrapped line."""
    gap = cur["y0"] - prev["y1"]
    if gap < -0.6 * cur["size"] or gap > 0.6 * cur["size"]:  # jumped up (new box/column) or paragraph spacing
        return True
    if slides:
        return False
    ends = prev["text"].rstrip().endswith((".", ":", "?", "!"))
    if ends and cur["x0"] > prev["x0"] + cur["size"]:  # indented first line (TeX)
        return True
    if prev["x1"] < 0.75 * right or (ends and prev["x1"] < 0.9 * right):  # short line: deliberate break
        return "line"
    return False


# ---------------------------------------------------------------- PPTX
def load_pptx(path):
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER

    prs = Presentation(path)
    H = prs.slide_height or 1
    skip_ph = {PP_PLACEHOLDER.FOOTER, PP_PLACEHOLDER.SLIDE_NUMBER, PP_PLACEHOLDER.DATE}
    title_ph = {PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE, PP_PLACEHOLDER.VERTICAL_TITLE}

    def flatten(shapes):
        for sh in shapes:
            if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
                yield from flatten(sh.shapes)
            else:
                yield sh

    def ph_type(sh):
        return sh.placeholder_format.type if sh.is_placeholder else None

    units = []
    for no, slide in enumerate(prs.slides, 1):
        shapes = [s for s in flatten(slide.shapes) if ph_type(s) not in skip_ph]
        images = sum(s.shape_type == MSO_SHAPE_TYPE.PICTURE for s in shapes)
        title_shape = next((s for s in shapes if ph_type(s) in title_ph and s.has_text_frame and s.text_frame.text.strip()), None)
        body = sorted((s for s in shapes if s is not title_shape), key=lambda s: ((s.top or 0), (s.left or 0)))
        title = _clean_ws(title_shape.text_frame.text) if title_shape else None
        lines = []
        if title is None and body and body[0].has_text_frame and (body[0].top or 0) < 0.25 * H:
            first = body[0].text_frame.text.strip().split("\n")[0].strip()
            if 0 < len(first) <= 80:  # untitled slide whose top text box acts as the title
                title = _clean_ws(first)
        if title:
            lines.append(line(title, "heading", 1, True))
        for sh in body:
            if getattr(sh, "has_table", False) and sh.has_table:
                rows = ["| " + " | ".join(_clean_ws(c.text) for c in r.cells) + " |" for r in sh.table.rows]
                lines.append(line("\n".join(rows), "table", 0, True))
            elif sh.has_text_frame:
                for para in sh.text_frame.paragraphs:
                    text = para.text.replace("\v", "\n").rstrip()
                    if not text.strip() or (title and _clean_ws(text) == title):
                        continue
                    fonts = [r.font.name for r in para.runs if r.text.strip()]
                    mono = bool(fonts) and all(f and MONO_FONT.search(f) for f in fonts)
                    if mono or looks_like_code(text):
                        lines.append(line(text, "code", para.level, False))
                    else:
                        lines.append(line(_clean_ws(text), "bullet", para.level, True))
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip() if slide.notes_slide.notes_text_frame else ""
            if notes:
                lines.append(line(notes, "notes", 0, True))
        units.append({"kind": "slide", "number": no, "title": title, "images": images, "lines": lines})

    master_text = []
    for m in prs.slide_masters:
        for shp in list(m.shapes) + [s for lay in m.slide_layouts for s in lay.shapes]:
            if shp.has_text_frame:
                t = _clean_ws(shp.text_frame.text)
                if t and not t.lower().startswith("click") and t not in master_text:
                    master_text.append(t)
    cp = prs.core_properties
    props = {k: v for k, v in {"title": cp.title, "author": cp.author, "last_modified_by": cp.last_modified_by}.items() if v}
    return {"layout": "slides", "units": units, "page_count": None, "slide_count": len(units),
            "properties": props, "master_text": master_text}


def _clean_ws(t):
    return re.sub(r"[ \t ]+", " ", t.replace("\v", " ")).strip()


# ---------------------------------------------------------------- PPT (binary PowerPoint 97-2003)
TITLE_TYPES = {0, 6}  # TextHeaderAtom: Title, CenterTitle
NOTES_TYPE = 2


def load_ppt(path):
    """Parse the 'PowerPoint Document' record stream directly ([MS-PPT]). Slide order comes from the
    SlideListWithText persist references resolved through the persist directory."""
    import olefile

    ole = olefile.OleFileIO(str(path))
    data = ole.openstream("PowerPoint Document").read()
    cur_user = ole.openstream("Current User").read()
    meta = ole.get_metadata()

    def records(off, end):
        while off + 8 <= end:
            vi, rt, rl = struct.unpack_from("<HHI", data, off)
            yield off, vi, rt, rl
            off += 8 + rl

    def texts(off, end, typ=None):
        for o, vi, rt, rl in records(off, end):
            if vi & 0xF == 0xF:
                yield from texts(o + 8, o + 8 + rl, typ)
            elif rt == 3999:  # TextHeaderAtom
                typ = struct.unpack_from("<I", data, o + 8)[0]
            elif rt == 4000:  # TextCharsAtom (UTF-16)
                yield typ, data[o + 8:o + 8 + rl].decode("utf-16-le", "replace")
            elif rt == 4008:  # TextBytesAtom (8-bit)
                yield typ, data[o + 8:o + 8 + rl].decode("latin-1")

    # persist directory: newest edit wins
    persist, edit, seen = {}, struct.unpack_from("<I", cur_user, 16)[0], set()
    doc_ref = None
    while edit and edit not in seen:
        seen.add(edit)
        last, pdir, docref = struct.unpack_from("<III", data, edit + 16)
        doc_ref = doc_ref or docref
        _, _, rl = struct.unpack_from("<HHI", data, pdir)
        o, end = pdir + 8, pdir + 8 + rl
        while o < end:
            entry = struct.unpack_from("<I", data, o)[0]
            o += 4
            for k in range(entry >> 20):
                persist.setdefault((entry & 0xFFFFF) + k, struct.unpack_from("<I", data, o)[0])
                o += 4
        edit = last
    doc_off = persist[doc_ref]
    _, _, doc_len = struct.unpack_from("<HHI", data, doc_off)
    order = []
    for o, vi, rt, rl in records(doc_off + 8, doc_off + 8 + doc_len):
        if rt == 4080 and vi >> 4 == 0:  # SlideListWithText, instance 0 = slides
            order = [struct.unpack_from("<I", data, o2 + 8)[0] for o2, _, r, _ in records(o + 8, o + 8 + rl) if r == 1011]

    units = []
    for no, pid in enumerate(order, 1):
        so = persist[pid]
        _, _, sl = struct.unpack_from("<HHI", data, so)
        title, lines, notes_ref = None, [], 0
        for o, vi, rt, rl in records(so + 8, so + 8 + sl):
            if rt == 1007:  # SlideAtom -> notesIdRef
                notes_ref = struct.unpack_from("<I", data, o + 8 + 16)[0]
        for typ, raw in texts(so + 8, so + 8 + sl):
            paras = [p.replace("\x0b", "\n").rstrip() for p in raw.split("\r")]
            if typ in TITLE_TYPES and title is None:
                title = _clean_ws(" ".join(paras))
                lines.append(line(title, "heading", 1, True))
                continue
            for p in paras:
                if not p.strip():
                    continue
                if looks_like_code(p):
                    lines.append(line(p.replace("\t", "    "), "code", 0, False))
                else:
                    lines.append(line(_clean_ws(p), "notes" if typ == NOTES_TYPE else "bullet", 0, True))
        if notes_ref and notes_ref in persist:
            no_off = persist[notes_ref]
            _, _, nl = struct.unpack_from("<HHI", data, no_off)
            notes = "\n".join(t for typ, t in texts(no_off + 8, no_off + 8 + nl) if typ == NOTES_TYPE).strip()
            if notes:
                lines.append(line(notes.replace("\r", "\n"), "notes", 0, True))
        units.append({"kind": "slide", "number": no, "title": title, "images": None, "lines": lines})
    props = {k: v.decode("latin-1") for k, v in {"title": meta.title, "author": meta.author,
                                                 "last_modified_by": meta.last_saved_by}.items() if v}
    master_text = [_clean_ws(t) for off in _master_offsets(data, doc_off, doc_len, records, persist)
                   for _, t in texts(off[0], off[1]) if t.strip() and not t.lower().startswith("click")]
    return {"layout": "slides", "units": units, "page_count": None, "slide_count": len(units),
            "properties": props, "master_text": list(dict.fromkeys(master_text))}


def _master_offsets(data, doc_off, doc_len, records, persist):
    for o, vi, rt, rl in records(doc_off + 8, doc_off + 8 + doc_len):
        if rt == 4080 and vi >> 4 == 1:  # SlideListWithText, instance 1 = masters
            for o2, _, r, _ in records(o + 8, o + 8 + rl):
                if r == 1011:
                    mo = persist.get(struct.unpack_from("<I", data, o2 + 8)[0])
                    if mo is not None:
                        yield mo + 8, mo + 8 + struct.unpack_from("<HHI", data, mo)[2]


# ---------------------------------------------------------------- DOCX (stdlib zip + XML)
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def load_docx(path):
    # ponytail: DOCX has no fixed pagination, so page numbers are null. Word's
    # <w:lastRenderedPageBreak/> markers could approximate them if a DOCX ever needs page citations.
    with zipfile.ZipFile(path) as z:
        root = ET.fromstring(z.read("word/document.xml"))
        core = ET.fromstring(z.read("docProps/core.xml")) if "docProps/core.xml" in z.namelist() else None
    lines = []
    for el in root.find(f"{W}body"):
        if el.tag == f"{W}p":
            style = el.find(f"{W}pPr/{W}pStyle")
            style = style.get(f"{W}val", "") if style is not None else ""
            text = "".join(t.text or "" if t.tag == f"{W}t" else ("\t" if t.tag == f"{W}tab" else "\n")
                           for t in el.iter() if t.tag in (f"{W}t", f"{W}tab", f"{W}br", f"{W}cr"))
            if not text.strip():
                continue
            fonts = [f.get(f"{W}ascii") for f in el.iter(f"{W}rFonts")]
            heading = re.match(r"(heading|title)\s*(\d*)", style, re.I)
            numbered = el.find(f"{W}pPr/{W}numPr")
            if heading:
                lines.append(line(_clean_ws(text), "heading", int(heading.group(2) or 1), True))
            elif re.search(r"code|source|html|verbatim", style, re.I) or (fonts and all(f and MONO_FONT.search(f) for f in fonts)) \
                    or looks_like_code(text):
                lines.append(line(text.rstrip(), "code", 0, False))
            elif numbered is not None or "list" in style.lower():
                ilvl = numbered.find(f"{W}ilvl") if numbered is not None else None
                lines.append(line(_clean_ws(text), "bullet", int(ilvl.get(f"{W}val")) if ilvl is not None else 0, True))
            else:
                lines.append(line(text.strip(), "text", 0, True))
        elif el.tag == f"{W}tbl":
            rows = []
            for tr in el.iter(f"{W}tr"):
                cells = ["".join(t.text or "" for t in tc.iter(f"{W}t")).strip() for tc in tr.iter(f"{W}tc")]
                rows.append("| " + " | ".join(cells) + " |")
            lines.append(line("\n".join(rows), "table", 0, True))
    props = {}
    if core is not None:
        for el in core:
            tag = el.tag.split("}")[1]
            if tag in ("title", "creator") and el.text:
                props["author" if tag == "creator" else tag] = el.text
    return {"layout": "document", "units": [{"kind": "section", "number": None, "title": None, "images": None, "lines": lines}],
            "page_count": None, "slide_count": None, "properties": props, "master_text": []}


# ---------------------------------------------------------------- TXT / MD
def load_text(path):
    raw = Path(path).read_text(encoding="utf-8", errors="replace")
    lines, in_code, brk = [], False, True
    for t in raw.splitlines():
        if t.strip().startswith("```"):
            in_code, brk = not in_code, True
            continue
        if in_code:
            lines.append(line(t, "code", 0, False))
        elif not t.strip():
            brk = True
        elif m := re.match(r"^(#{1,6})\s+(.*)", t):
            lines.append(line(m.group(2).strip(), "heading", len(m.group(1)), True))
            brk = True
        else:
            lines.append(line(t.rstrip(), "code" if looks_like_code(t) else "text", 0, brk))
            brk = False
    return {"layout": "document", "units": [{"kind": "section", "number": None, "title": None, "images": None, "lines": lines}],
            "page_count": None, "slide_count": None, "properties": {}, "master_text": []}


LOADERS = {"pdf": load_pdf, "pptx": load_pptx, "ppt": load_ppt, "docx": load_docx, "txt": load_text, "md": load_text}


def extract(path, file_type):
    """Dispatch to the format loader. Raises on failure; the pipeline records the error."""
    return LOADERS[file_type](path)
