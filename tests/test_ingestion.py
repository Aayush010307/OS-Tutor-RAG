"""Phase 1 ingestion tests. Fixtures are generated at runtime so no binary files live in the repo.

    python -m pytest tests -q
"""
import hashlib
import json
import shutil
import zipfile
from pathlib import Path

import fitz
import pytest
from pptx import Presentation
from pptx.util import Inches, Pt

from src.ingestion import chunker, cleaner, loaders, pipeline, scanner
from src.ingestion.dedup import find_duplicates

ROOT = Path(__file__).resolve().parents[1]
REAL_PPT = ROOT / "Docs/Synchronisation/FALLSEM2025-26_VL_BCSE303L_00100_TH_2025-09-11_process-Synchronization_locks_semaphores_monitors.ppt"


# ------------------------------------------------------------------ fixture builders
def make_notes_pdf(path):
    """3-page textbook-style PDF: running header, page numbers, heading, prose and a code listing."""
    doc = fitz.open()
    for n in range(1, 4):
        p = doc.new_page()
        p.insert_text((72, 30), "OPERATING SYSTEMS NOTES", fontsize=8)
        p.insert_text((300, 820), str(n), fontsize=9)
        y = 90
        if n == 1:
            p.insert_text((72, y), "Condition Variables", fontsize=18, fontname="hebo")
            y += 40
        for i in range(12):
            p.insert_text((72, y), f"Page {n} line {i}: a thread waits on a condition variable until another thread signals it.", fontsize=11)
            y += 14
        if n == 2:
            y += 20
            for code in ["pthread_mutex_lock(&m);", "while (done == 0)", "    pthread_cond_wait(&c, &m);", "pthread_mutex_unlock(&m);"]:
                p.insert_text((72, y), code, fontsize=11, fontname="cour")
                y += 14
    doc.save(path)


def make_qa_pdf(path):
    """Practice problems where question 2's answer continues on page 2."""
    doc = fitz.open()
    p = doc.new_page()
    y = 60
    for text, size, font in [("Lectures on Operating Systems (Jane Doe, IIT Bombay)", 13, "helv"),
                             ("Practice Problems: Synchronization", 20, "hebo")]:
        p.insert_text((72, y), text, fontsize=size, fontname=font)
        y += 40
    lines = ["1. What is a race condition?", "Ans: The outcome depends on the timing of thread interleavings.", "",
             "2. Explain why pthread_cond_wait() must be called with the mutex held."]
    for t in lines:
        if t:
            p.insert_text((72, y), t, fontsize=11)
        y += 22
    p2 = doc.new_page()
    y = 60
    for t in ["Ans: Otherwise a signal can arrive between checking the condition and sleeping,", "causing a missed wakeup.", "",
              "3. Is a spinlock useful on a single-core machine? [T/F]", "Ans: F", "",
              "4. Repeat the previous question for a multi-core machine.", "Ans: T"]:
        if t:
            p2.insert_text((72, y), t, fontsize=11)
        y += 22
    doc.save(path)


def make_pptx(path):
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = "Semaphores"
    tf = s.placeholders[1].text_frame
    tf.text = "A semaphore is an integer with atomic operations"
    for text, level in [("sem_wait() decrements and may block", 1), ("sem_post() increments and wakes a waiter", 1)]:
        para = tf.add_paragraph()
        para.text, para.level = text, level
    s.notes_slide.notes_text_frame.text = "Mention the binary semaphore as a lock."
    s2 = prs.slides.add_slide(prs.slide_layouts[5])
    s2.shapes.title.text = "Producer code"
    box = s2.shapes.add_textbox(Inches(1), Inches(2), Inches(6), Inches(3)).text_frame
    for i, code in enumerate(["sem_wait(&empty);", "sem_wait(&mutex);", "put(item);", "sem_post(&mutex);", "sem_post(&full);"]):
        para = box.paragraphs[0] if i == 0 else box.add_paragraph()
        para.text = code
        para.runs[0].font.name = "Courier New"
        para.runs[0].font.size = Pt(14)
    prs.save(path)


def make_docx(path):
    W = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
    body = (
        '<w:p><w:pPr><w:pStyle w:val="Heading1"/></w:pPr><w:r><w:t>Deadlock</w:t></w:r></w:p>'
        '<w:p><w:r><w:t>Deadlock needs mutual exclusion, hold and wait, no preemption and circular wait.</w:t></w:r></w:p>'
        '<w:p><w:pPr><w:pStyle w:val="ListParagraph"/><w:numPr><w:ilvl w:val="0"/><w:numId w:val="1"/></w:numPr></w:pPr>'
        '<w:r><w:t>Break circular wait by ordering locks</w:t></w:r></w:p>'
        '<w:tbl><w:tr><w:tc><w:p><w:r><w:t>Condition</w:t></w:r></w:p></w:tc><w:tc><w:p><w:r><w:t>Prevention</w:t></w:r></w:p></w:tc></w:tr>'
        '<w:tr><w:tc><w:p><w:r><w:t>Circular wait</w:t></w:r></w:p></w:tc><w:tc><w:p><w:r><w:t>Lock ordering</w:t></w:r></w:p></w:tc></w:tr></w:tbl>'
    )
    with zipfile.ZipFile(path, "w") as z:
        z.writestr("[Content_Types].xml", "<Types/>")
        z.writestr("word/document.xml", f'<?xml version="1.0"?><w:document {W}><w:body>{body}</w:body></w:document>')


@pytest.fixture(scope="session")
def corpus(tmp_path_factory):
    root = tmp_path_factory.mktemp("proj")
    src = root / "Docs"
    (src / "Threads").mkdir(parents=True)
    (src / "Synchronisation").mkdir()
    make_notes_pdf(src / "Threads" / "cv-notes.pdf")
    make_qa_pdf(src / "Synchronisation" / "ps-sync.pdf")
    make_pptx(src / "Synchronisation" / "FALLSEM2025-26_VL_BCSE303L_00100_TH_2025-09-16_Semaphores.pptx")
    shutil.copy(src / "Synchronisation" / "FALLSEM2025-26_VL_BCSE303L_00100_TH_2025-09-16_Semaphores.pptx",
                src / "Synchronisation" / "Semaphores - Copy.pptx")
    make_docx(src / "Synchronisation" / "deadlock.docx")
    (src / "Threads" / "notes.md").write_text("# Thread API\n\nUse pthread_join() to wait.\n\n```\npthread_join(t, NULL);\n```\n")
    (src / "Threads" / "scan.djvu").write_bytes(b"not supported")
    before = {p: hashlib.sha256(p.read_bytes()).hexdigest() for p in src.rglob("*") if p.is_file()}
    assert pipeline.run(src, root / "data") == 0
    return {"root": root, "src": src, "out": root / "data", "before": before,
            "docs": {d["filename"]: d for d in json.loads((root / "data/manifests/documents.json").read_text())},
            "chunks": [json.loads(l) for l in (root / "data/chunks/chunks.jsonl").read_text().splitlines()]}


def chunks_of(corpus, filename):
    did = corpus["docs"][filename]["document_id"]
    return [c for c in corpus["chunks"] if c["document_id"] == did]


# ------------------------------------------------------------------ extraction
def test_pdf_extraction_pages_heading_code_and_header_removal(corpus):
    ext = loaders.extract(corpus["src"] / "Threads" / "cv-notes.pdf", "pdf")
    assert [u["number"] for u in ext["units"]] == [1, 2, 3]
    cleaned, stats = cleaner.clean(ext)
    blocks = [b for u in cleaned["units"] for b in u["blocks"]]
    assert blocks[0] == {"type": "heading", "text": "Condition Variables", "level": 1}
    assert stats["margin_lines_removed"] >= 6  # running header + page number on each page
    assert not any("OPERATING SYSTEMS NOTES" in b["text"] for b in blocks)
    code = [b for b in cleaned["units"][1]["blocks"] if b["type"] == "code"]
    assert code and code[0]["text"].splitlines() == ["pthread_mutex_lock(&m);", "while (done == 0)",
                                                     "    pthread_cond_wait(&c, &m);", "pthread_mutex_unlock(&m);"]


def test_pptx_extraction_slides_titles_bullets_notes():
    import tempfile
    with tempfile.TemporaryDirectory() as d:
        make_pptx(Path(d) / "s.pptx")
        ext = loaders.extract(Path(d) / "s.pptx", "pptx")
    assert ext["slide_count"] == 2 and [u["number"] for u in ext["units"]] == [1, 2]
    assert [u["title"] for u in ext["units"]] == ["Semaphores", "Producer code"]
    lines = ext["units"][0]["lines"]
    assert [(l["kind"], l["level"]) for l in lines if l["kind"] == "bullet"] == [("bullet", 0), ("bullet", 1), ("bullet", 1)]
    assert any(l["kind"] == "notes" and "binary semaphore" in l["text"] for l in lines)
    assert all(l["kind"] == "code" for l in ext["units"][1]["lines"][1:])


def test_docx_extraction_heading_list_table(corpus):
    ext = loaders.extract(corpus["src"] / "Synchronisation" / "deadlock.docx", "docx")
    kinds = [(l["kind"], l["text"]) for l in ext["units"][0]["lines"]]
    assert kinds[0] == ("heading", "Deadlock")
    assert ("bullet", "Break circular wait by ordering locks") in kinds
    assert ("table", "| Condition | Prevention |\n| Circular wait | Lock ordering |") in kinds


@pytest.mark.skipif(not REAL_PPT.exists(), reason="real .ppt corpus file not present")
def test_legacy_ppt_extraction_on_real_file():
    ext = loaders.extract(REAL_PPT, "ppt")
    assert ext["slide_count"] == 59
    assert ext["units"][0]["title"] == "Process Synchronization"
    assert ext["units"][3]["title"] == "Producer-Consumer Problem"


def test_unsupported_and_hidden_files_are_reported(corpus):
    manifest = json.loads((corpus["out"] / "manifests/corpus_manifest.json").read_text())
    assert [u["relative_path"] for u in manifest["unsupported_files"]] == ["Threads/scan.djvu"]
    assert manifest["totals"]["failed"] == 0


# ------------------------------------------------------------------ metadata
def test_metadata_is_evidence_based(corpus):
    ps = corpus["docs"]["ps-sync.pdf"]
    assert (ps["source"], ps["professor"], ps["document_type"], ps["topic_area"]) == \
           ("IIT Bombay", "Jane Doe", "practice_problems", "Synchronisation")
    vit = corpus["docs"]["FALLSEM2025-26_VL_BCSE303L_00100_TH_2025-09-16_Semaphores.pptx"]
    assert (vit["source"], vit["semester"], vit["course_code"], vit["document_type"]) == ("VIT", "FALLSEM2025-26", "BCSE303L", "lecture")
    assert vit["professor"] is None  # never taken from embedded file properties
    notes = corpus["docs"]["cv-notes.pdf"]
    assert (notes["source"], notes["professor"], notes["module"], notes["topic_area"]) == ("unknown", None, None, "Threads")
    assert "Condition Variables" in notes["subtopic"]
    assert "True/False" in ps["problem_types"]


# ------------------------------------------------------------------ determinism, dedup
def test_deterministic_ids_and_chunks(corpus, tmp_path):
    assert pipeline.run(corpus["src"], tmp_path / "again") == 0
    again = (tmp_path / "again/chunks/chunks.jsonl").read_text()
    assert again == (corpus["out"] / "chunks/chunks.jsonl").read_text()
    assert scanner.make_doc_id("Threads/cv-notes.pdf") == scanner.make_doc_id("Threads/cv-notes.pdf")
    assert all(c["chunk_id"] == f"{c['document_id']}::c{c['chunk_index']:04d}" for c in corpus["chunks"])


def test_duplicate_detection(corpus):
    copy = corpus["docs"]["Semaphores - Copy.pptx"]
    orig = corpus["docs"]["FALLSEM2025-26_VL_BCSE303L_00100_TH_2025-09-16_Semaphores.pptx"]
    assert (copy["is_duplicate"], copy["duplicate_of"], copy["duplicate_type"]) == (True, orig["document_id"], "exact_file")
    assert not orig["is_duplicate"]
    assert all(c["duplicate_of"] and c["is_duplicate_document"] for c in chunks_of(corpus, "Semaphores - Copy.pptx"))
    assert (corpus["src"] / "Synchronisation" / "Semaphores - Copy.pptx").exists()  # recorded, never deleted


def test_near_duplicate_text_is_recorded_not_merged():
    base = " ".join(f"word{i}" for i in range(300))
    docs = [{"document_id": "a", "filename": "a.pdf", "file_hash": "1", "text": base},
            {"document_id": "b", "filename": "b.pdf", "file_hash": "2", "text": base + " extra tail words here"}]
    out = find_duplicates(docs)
    assert not out["b"]["is_duplicate"]
    assert out["a"]["near_duplicates"][0]["document_id"] == "b"


# ------------------------------------------------------------------ chunking
def _doc(dtype="notes"):
    return {"document_id": "d", "filename": "d.pdf", "relative_path": "Docs/d.pdf", "source": "unknown",
            "document_type": dtype, "topic_area": "Threads"}


def _cleaned(blocks, per_page=1, layout="document"):
    units = [{"kind": "page", "number": i // per_page + 1, "title": None, "images": 0, "blocks": []} for i in range(0, len(blocks), per_page)]
    for i, b in enumerate(blocks):
        units[i // per_page]["blocks"].append(b)
    return {"layout": layout, "units": units}


def test_chunking_splits_large_sections_and_merges_tiny_ones():
    para = "Threads share the address space of their process but each has its own stack and registers. " * 6
    blocks = [{"type": "heading", "text": "Tiny intro", "level": 1},
              {"type": "heading", "text": "Big section", "level": 2}] + [{"type": "paragraph", "text": para, "level": 0}] * 12
    chunks = chunker.chunk_document(_doc(), _cleaned(blocks, per_page=3))
    assert len(chunks) > 2
    assert all(c["token_count"] <= chunker.MAX_TOKENS for c in chunks)
    assert chunks[0]["text"].startswith("# Tiny intro\n\n## Big section")  # tiny heading merged forward
    assert all(c["section"] == "Tiny intro > Big section" for c in chunks)
    assert [c["page_start"] for c in chunks] == sorted(c["page_start"] for c in chunks)


def test_code_preservation_and_content_type():
    code = "\n".join(f"pthread_mutex_lock(&m{i});\ncounter{i}++;\npthread_mutex_unlock(&m{i});" for i in range(120))
    blocks = [{"type": "heading", "text": "Locks", "level": 1}, {"type": "paragraph", "text": "Example:", "level": 0},
              {"type": "code", "text": code, "level": 0}]
    chunks = chunker.chunk_document(_doc(), _cleaned(blocks, per_page=3))
    assert all(c["token_count"] <= chunker.MAX_TOKENS for c in chunks)
    joined = "\n".join(c["text"] for c in chunks)
    assert "pthread_mutex_lock(&m0);\ncounter0++;\npthread_mutex_unlock(&m0);" in joined  # lines not merged into prose
    assert all(c["content_type"] in ("code", "mixed") for c in chunks)
    assert all(c["text"].count("```") % 2 == 0 for c in chunks)  # every code piece is fenced


def test_question_answer_pairs_stay_together_across_pages(corpus):
    chunks = chunks_of(corpus, "ps-sync.pdf")
    by_q = {c["question_number"]: c for c in chunks}
    assert set(by_q) == {1, 2, 3, 4}
    q2 = by_q[2]
    assert "pthread_cond_wait()" in q2["text"] and "missed wakeup" in q2["text"]
    assert (q2["page_start"], q2["page_end"], q2["has_answer"], q2["chunk_type"]) == (1, 2, True, "qa_pair")
    assert "Ans: F" in by_q[3]["text"] and "True/False" in by_q[3]["problem_types"]
    assert by_q[4]["refers_to_question"] == 3


def test_every_chunk_is_traceable(corpus):
    for c in corpus["chunks"]:
        assert c["text"].strip() and c["document_id"] and c["filename"] and c["relative_path"]
        ft = corpus["docs"][c["filename"]]["file_type"]
        if ft == "pdf":
            assert c["page_start"] and c["page_end"] >= c["page_start"]
        if ft in ("pptx", "ppt"):
            assert c["slide_start"] and c["slide_end"] >= c["slide_start"]
    slides = chunks_of(corpus, "FALLSEM2025-26_VL_BCSE303L_00100_TH_2025-09-16_Semaphores.pptx")
    assert slides[0]["section"] == "Semaphores" and slides[0]["slide_start"] == 1
    issues = json.loads((corpus["out"] / "manifests/validation.json").read_text())
    assert not [i for i in issues if i["level"] == "error"]


def test_raw_corpus_is_untouched(corpus):
    after = {p: hashlib.sha256(p.read_bytes()).hexdigest() for p in corpus["src"].rglob("*") if p.is_file()}
    assert after == corpus["before"]


def test_refuses_output_inside_source(corpus):
    with pytest.raises(SystemExit):
        pipeline.run(corpus["src"], corpus["src"] / "data")
